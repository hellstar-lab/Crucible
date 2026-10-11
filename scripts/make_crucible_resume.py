"""Build tests/fixtures/docs/docx_crucible_resume.docx: a resume that hellstar-lab/Crucible can actually answer (Phase 4K K5).

    python scripts/make_crucible_resume.py

Every line is chosen against the repository's default branch (main @ 2bffef4), measured with git blame:
  * skills: libraries requirements.txt lists AND the code imports (FastAPI is listed but never imported: a real gap);
  * "Built the blame attribution module ...": backend/evidence/blame.py, 184/184 lines by hellstar-lab -> supported;
  * "Designed the unicode scanning module ...": backend/shield/unicode_scan.py, 164/164 lines by another account
    (Auratech350) -> the contradiction path, which needs Super 3/3 AND Ultra 3/3;
  * "Reduced PDF analysis latency by 40%": a metric with no benchmark in the repository -> no_evidence;
  * "Led a team of 4 engineers": team size -> unverifiable;
  * the contact line links github.com/hellstar-lab and the repository, so the identity check agrees.
Not part of the frozen 2C ground truth (MANIFEST.json): a demo input, rebuilt deterministically by this script.
"""

from pathlib import Path

import docx

OUT = Path(__file__).resolve().parent.parent / "tests" / "fixtures" / "docs" / "docx_crucible_resume.docx"


def main() -> None:
    d = docx.Document()
    d.core_properties.author = "Crucible demo"
    d.core_properties.title = "Demo resume for hellstar-lab"
    d.add_heading("Crucible Demo Candidate", 0)
    d.add_paragraph("github.com/hellstar-lab | github.com/hellstar-lab/Crucible")
    d.add_heading("Projects", 1)
    d.add_heading("Crucible, hiring-evidence system", 3)
    for line in ("Built the blame attribution module that maps every source line to its author.",
                 "Designed the unicode scanning module for hidden-character detection.",
                 "Reduced PDF analysis latency by 40%.",
                 "Led a team of 4 engineers."):
        d.add_paragraph(line, style="List Bullet")
    d.add_heading("Skills", 1)
    d.add_paragraph("Python, pytest, Pydantic, tree-sitter, FastAPI, httpx, pdfplumber, pypdf, python-docx, RapidFuzz, PyYAML")
    d.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()

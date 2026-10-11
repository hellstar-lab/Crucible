"""Phase 4J J1: skills lists decompose deterministically, with no model call. Offline (conftest removes the key)."""

from __future__ import annotations

import logging
from pathlib import Path

import pytest

from backend.claims import pipeline, skills
from backend.retrieval import retrieve
from backend.shield.analyze import analyze_document

from test_claims_pipeline import _UNIT, FakeLLM  # noqa: F401  (same scripted fake model)

DOCS = Path(__file__).resolve().parent / "fixtures" / "docs"


@pytest.fixture
def fake(monkeypatch):
    from backend import llm
    f = FakeLLM()
    monkeypatch.setattr(llm, "call", f)
    return f


def _skills(out):
    return [c for c in out.claims if c.decomposition_tier == "deterministic"]


def _skills_prompts(fake, words=("Python", "PostgreSQL", "Kubernetes", "Docker")):
    """Prompts whose UNIT is a skills line (a skills line may still appear in another unit's context)."""
    return [p for _t, _tier, p in fake.prompts if any(w in _UNIT.search(p).group(1) for w in words)]


def test_clean_resume_skills_line_becomes_one_bare_claim_per_item(fake):
    logging.disable(logging.WARNING)
    out = pipeline.extract_claims(analyze_document(DOCS / "docx_clean_resume.docx"), decomposition_tier="nano")
    got = _skills(out)
    assert [c.text for c in got] == ["Python", "Go", "PostgreSQL", "Docker", "Kubernetes"]
    for c in got:
        assert (c.kind, c.scope, c.checkable_from_code, c.qualifiers) == ("technology_use", "none", "yes", [])
        assert out.document_text[c.source_span[0]:c.source_span[1]] == c.text  # the span is the item itself
    assert not _skills_prompts(fake) and out.counters["skills_claims"] == 5
    assert not any("proficient" in c.text.lower() or "experience" in c.text.lower() for c in got)


def test_hidden_text_resume_label_line_is_split_and_the_label_stripped(fake):
    out = pipeline.extract_claims(analyze_document(DOCS / "docx_hidden.docx"), decomposition_tier="nano")
    assert [c.text for c in _skills(out)] == ["Python", "Docker"]  # "Skills:" stripped, trailing "." dropped
    units_sent = [_UNIT.search(p).group(1) for _t, _tier, p in fake.prompts]
    assert units_sent and not any("Python" in u for u in units_sent)  # it may appear as a neighbour's context, never as a unit


def test_skills_only_resume(fake, tmp_path):
    """Stands in for c02_skills_only (no such fixture exists in this repo): a resume that is nothing but skills."""
    import docx
    d = docx.Document()
    d.add_heading("Sam Rivera", 0)
    d.add_heading("Technical Skills", 1)
    d.add_paragraph("Languages: Python, Go; TypeScript | Rust")
    d.add_paragraph("Frameworks: FastAPI, React (Redux, Hooks), Django")
    d.add_paragraph("Cloud & DevOps: Docker • Kubernetes • Terraform, and AWS")
    d.add_paragraph("Data: PostgreSQL, Redis, Snowflake")
    path = tmp_path / "skills_only.docx"
    d.save(path)
    out = pipeline.extract_claims(analyze_document(path), decomposition_tier="nano")
    assert [c.text for c in _skills(out)] == ["Python", "Go", "TypeScript", "Rust", "FastAPI", "React (Redux, Hooks)", "Django",
                                              "Docker", "Kubernetes", "Terraform", "AWS", "PostgreSQL", "Redis", "Snowflake"]
    assert fake.prompts == []  # nothing in this resume reached a model
    for c in _skills(out):
        assert " ".join(out.document_text[c.source_span[0]:c.source_span[1]].split()) == c.text


@pytest.mark.parametrize("text,expected", [
    ("Python, Go, PostgreSQL", ["Python", "Go", "PostgreSQL"]),
    ("Languages: Python, Go", ["Python", "Go"]),
    ("Python; Go | Rust • C++", ["Python", "Go", "Rust", "C++"]),
    ("Python (Django, Flask), Node.js, CI/CD", ["Python (Django, Flask)", "Node.js", "CI/CD"]),
    ("Python, Go, and Rust.", ["Python", "Go", "Rust"]),
])
def test_split_items(text, expected):
    items = skills.split_items(text, 0, len(text))
    assert [t for _a, _b, t in items] == expected
    assert all(text[a:b] == t for a, b, t in items)


def test_prose_under_a_skills_heading_is_left_to_the_model():
    t = "Strong communication skills developed while presenting quarterly results to the leadership team"
    assert skills.split_items(t, 0, len(t)) is None


def test_cap_records_the_overflow(fake):
    items = ", ".join(f"Tool{i}" for i in range(45))
    out = pipeline.extract_from_text(items, headings=["SKILLS"], decomposition_tier="nano")
    assert len(_skills(out)) == skills.MAX_ITEMS and out.counters["skills_items_over_cap"] == 5
    over = [e for e in out.excluded if e.reason_code == "list_cap"]
    assert len(over) == 1 and "Tool44" in over[0].detail and fake.prompts == []


@pytest.mark.parametrize("heading,unit,is_skills", [
    ("Skills", "Python, Go", True), ("Tech Stack", "Go, gRPC", True), ("Experience", "Skills: Python, Docker.", True),
    ("Experience", "Built a caching layer with Redis", False), (None, "Python, Go", False),
])
def test_is_skills_unit(heading, unit, is_skills):
    assert skills.is_skills_unit(unit, heading, [heading] if heading else []) is is_skills


def test_unknown_technology_is_reported_not_dropped():
    es = retrieve({"claim_id": "s.c0", "kind": "technology_use", "checkable_from_code": "yes", "text": "Snowflake"}, [])
    assert not es.retrieval_performed and es.skip_reason.startswith("unknown_technology: 'Snowflake'")
